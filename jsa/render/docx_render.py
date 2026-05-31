"""DocxRenderer: Markdown -> ATS-friendly DOCX via python-docx."""

import asyncio
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_LINE_SPACING

from jsa.render.base import Renderer


def _add_bottom_border(paragraph) -> None:
    """Add a thin bottom border to the given paragraph using OOXML manipulation."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    for k, v in (
        ("w:val", "single"),
        ("w:sz", "6"),
        ("w:space", "1"),
        ("w:color", "auto"),
    ):
        bottom.set(qn(k), v)
    pBdr.append(bottom)


def _add_runs(paragraph, text: str, default_bold: bool = False) -> None:
    """Add runs to a paragraph, respecting **bold** markers.

    If default_bold is True, the entire run is bold regardless of markers
    (used for headings that are inherently bold).
    """
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        if default_bold:
            run.bold = True
        else:
            # Odd-indexed parts are between ** markers → bold
            run.bold = (i % 2 == 1)


class DocxRenderer(Renderer):
    name = "docx"

    async def render(self, markdown: str, output_path: Path) -> None:
        await asyncio.to_thread(self._render_sync, markdown, output_path)

    def _render_sync(self, markdown: str, output_path: Path) -> None:
        document = Document()

        # --- Page margins: 0.75 inches all sides ---
        section = document.sections[0]
        margin = Inches(0.75)
        section.top_margin = margin
        section.bottom_margin = margin
        section.left_margin = margin
        section.right_margin = margin

        # --- Default body style: Calibri 10.5pt, 1.3 line spacing ---
        normal_style = document.styles["Normal"]
        normal_font = normal_style.font
        normal_font.name = "Calibri"
        normal_font.size = Pt(10.5)
        # Remove any color on Normal (ensures no coloring bleeds through)
        normal_font.color.rgb = RGBColor(0, 0, 0)
        normal_pf = normal_style.paragraph_format
        normal_pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        normal_pf.line_spacing = 1.3

        # Parse lines
        lines = markdown.splitlines()

        found_h1 = False
        expect_contact = False  # True immediately after H1 is processed
        last_paragraph = None   # Track for bottom-border on ---

        for line in lines:
            stripped = line.strip()

            # Skip blank lines
            if not stripped:
                if expect_contact:
                    # Blank line between H1 and contact — keep waiting
                    pass
                continue

            # H1 — name heading
            if stripped.startswith("# ") and not stripped.startswith("## "):
                name_text = stripped[2:].strip()
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                pf.line_spacing = 1.3
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                run = p.add_run(name_text)
                run.bold = True
                run.font.name = "Calibri"
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(0, 0, 0)
                last_paragraph = p
                found_h1 = True
                expect_contact = True
                continue

            # Contact line — the first non-blank line after H1
            if expect_contact:
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                pf.line_spacing = 1.3
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                run = p.add_run(stripped)
                run.italic = True
                run.font.name = "Calibri"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0, 0, 0)
                last_paragraph = p
                expect_contact = False
                continue

            # Horizontal rule — add bottom border to preceding paragraph, no blank paragraph
            if stripped in ("---", "***"):
                if last_paragraph is not None:
                    _add_bottom_border(last_paragraph)
                # Don't add a new paragraph; the border goes on the last one
                continue

            # H2 — section heading
            if stripped.startswith("## "):
                heading_text = stripped[3:].strip()
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf = p.paragraph_format
                pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                pf.line_spacing = 1.3
                pf.space_before = Pt(2)
                pf.space_after = Pt(0)
                run = p.add_run(heading_text)
                run.bold = True
                run.font.name = "Calibri"
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(0, 0, 0)
                last_paragraph = p
                continue

            # H3 — sub-heading, treated as body bold
            if stripped.startswith("### "):
                heading_text = stripped[4:].strip()
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf = p.paragraph_format
                pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                pf.line_spacing = 1.3
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                _add_runs(p, heading_text, default_bold=True)
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(10.5)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                last_paragraph = p
                continue

            # Bullet point
            if stripped.startswith("- ") or stripped.startswith("* "):
                bullet_text = stripped[2:].strip()
                p = document.add_paragraph(style="List Bullet")
                pf = p.paragraph_format
                pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                pf.line_spacing = 1.3
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                _add_runs(p, bullet_text)
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(10.5)
                    if run.font.color.type is not None:
                        run.font.color.rgb = RGBColor(0, 0, 0)
                last_paragraph = p
                continue

            # Body paragraph
            p = document.add_paragraph()
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            pf.line_spacing = 1.3
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            _add_runs(p, stripped)
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(10.5)
                run.font.color.rgb = RGBColor(0, 0, 0)
            last_paragraph = p

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))
